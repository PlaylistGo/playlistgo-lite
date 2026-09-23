from openpyxl import Workbook, load_workbook
import os

from pathlib import Path

FILE = Path(__file__).parent / "songs.xlsx"
SHEET = "Active"


def makesure_workbook():
    if os.path.exists(FILE):
        return
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET
    sheet.append(["Title", "YouTube Link"])
    workbook.save(FILE)


def read_playlist():
    makesure_workbook()
    workbook = load_workbook(FILE)
    sheet = workbook[SHEET]

    songs = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        title, link = row[0], row[1]
        if title:
            songs.append((title, link))
    return songs


def write_playlist(songs):
    workbook = load_workbook(FILE)
    sheet = workbook[SHEET]

    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)

    for i, (title, link) in enumerate(songs, start=2):
        sheet.cell(row=i, column=1, value=title)
        sheet.cell(row=i, column=2, value=link)

    workbook.save(FILE)
