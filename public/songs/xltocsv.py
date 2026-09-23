import csv
from openpyxl import load_workbook

SRC = "songs.xlsx"
SHEET = "Active"
OUT = "songs.csv"

wb = load_workbook(SRC, read_only=True)
sheet = wb[SHEET]

rows = sheet.iter_rows(values_only=True)
headers = [str(h).strip() if h else "" for h in next(rows)]
print("Headers found:", headers)

title_idx = headers.index("Title")
url_idx = headers.index("YouTube Link")

with open(OUT, "w", newline="", encoding="utf-8") as f:
    writer = csv.writer(f)
    writer.writerow(["title", "url"])
    for row in rows:
        title = row[title_idx] or "Untitled"
        url = row[url_idx] or ""
        writer.writerow([title, url])

print(f"Wrote {OUT}")